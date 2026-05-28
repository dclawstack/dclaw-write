"""Document exports — Markdown / HTML / DOCX.

Markdown is stored as-is, HTML is a minimal converter that handles headings,
lists, paragraphs, and bold/italic. DOCX uses ``python-docx``.

Kept dependency-light: no `markdown` lib, no `mammoth`, no `pandoc`.
"""
from __future__ import annotations

import html
import io
import re

from docx import Document as DocxDocument
from docx.shared import Pt

_BOLD = re.compile(r"\*\*(.+?)\*\*")
_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
_INLINE_CODE = re.compile(r"`([^`]+)`")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_ULIST = re.compile(r"^[\-\*]\s+(.*)$")
_OLIST = re.compile(r"^\d+\.\s+(.*)$")


def to_markdown(title: str, content: str) -> str:
    header = f"# {title.strip()}\n\n" if title.strip() else ""
    return f"{header}{content.rstrip()}\n"


def to_html(title: str, content: str) -> str:
    body_parts: list[str] = []
    list_buffer: list[str] = []
    list_type: str | None = None

    def flush_list() -> None:
        nonlocal list_buffer, list_type
        if list_type and list_buffer:
            items = "".join(f"<li>{_inline(item)}</li>" for item in list_buffer)
            body_parts.append(f"<{list_type}>{items}</{list_type}>")
        list_buffer = []
        list_type = None

    for raw_line in content.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            flush_list()
            continue
        heading = _HEADING.match(line)
        if heading:
            flush_list()
            level = len(heading.group(1))
            body_parts.append(f"<h{level}>{_inline(heading.group(2))}</h{level}>")
            continue
        ulist = _ULIST.match(line)
        olist = _OLIST.match(line)
        if ulist:
            if list_type != "ul":
                flush_list()
                list_type = "ul"
            list_buffer.append(ulist.group(1))
            continue
        if olist:
            if list_type != "ol":
                flush_list()
                list_type = "ol"
            list_buffer.append(olist.group(1))
            continue
        flush_list()
        body_parts.append(f"<p>{_inline(line)}</p>")
    flush_list()

    safe_title = html.escape(title.strip() or "Untitled")
    body = "\n".join(body_parts) or "<p></p>"
    return (
        "<!doctype html>\n<html lang=\"en\">\n<head>\n"
        f"  <meta charset=\"utf-8\">\n  <title>{safe_title}</title>\n"
        "  <style>body{font-family:Georgia,serif;max-width:42rem;margin:2rem auto;"
        "padding:0 1rem;line-height:1.6;color:#1f2937}h1,h2,h3{font-family:"
        "-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;color:#0f172a}"
        "</style>\n</head>\n<body>\n"
        f"  <h1>{safe_title}</h1>\n  {body}\n</body>\n</html>\n"
    )


def to_docx(title: str, content: str) -> bytes:
    doc = DocxDocument()
    styles = doc.styles["Normal"]
    styles.font.name = "Georgia"
    styles.font.size = Pt(11)

    if title.strip():
        doc.add_heading(title.strip(), level=0)

    for paragraph in _paragraph_blocks(content):
        first_line = paragraph.splitlines()[0]
        heading = _HEADING.match(first_line)
        if heading and len(paragraph.splitlines()) == 1:
            level = min(len(heading.group(1)), 4)
            doc.add_heading(heading.group(2), level=level)
            continue
        if all(_ULIST.match(l) or _OLIST.match(l) for l in paragraph.splitlines()):
            for line in paragraph.splitlines():
                match = _ULIST.match(line) or _OLIST.match(line)
                doc.add_paragraph(match.group(1), style="List Bullet")
            continue
        doc.add_paragraph(paragraph)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paragraph_blocks(content: str) -> list[str]:
    blocks: list[str] = []
    buf: list[str] = []
    for line in content.splitlines():
        if line.strip():
            buf.append(line)
        elif buf:
            blocks.append("\n".join(buf))
            buf = []
    if buf:
        blocks.append("\n".join(buf))
    return blocks


def _inline(text: str) -> str:
    escaped = html.escape(text)
    escaped = _BOLD.sub(r"<strong>\1</strong>", escaped)
    escaped = _ITALIC.sub(r"<em>\1</em>", escaped)
    escaped = _INLINE_CODE.sub(r"<code>\1</code>", escaped)
    return escaped


def safe_filename(title: str, fallback: str = "document") -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", title.strip()).strip("-")
    return slug or fallback
